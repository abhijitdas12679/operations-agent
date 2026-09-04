import os
import re
from docx import Document


OUTPUT_DIR = "outputs/templates"


def _safe_filename(name: str) -> str:
    name = re.sub(r'[\\/*?:"<>|]', "", name or "template")
    name = name.strip().replace(" ", "_")
    return name[:120] or "template"


def create_template_docx(template) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    filename = f"{_safe_filename(template.template_name)}.docx"
    file_path = os.path.join(OUTPUT_DIR, filename)

    doc = Document()

    doc.add_heading(template.template_name, level=1)

    if template.category:
        p = doc.add_paragraph()
        p.add_run("Category: ").bold = True
        p.add_run(template.category)

    if template.template_type:
        p = doc.add_paragraph()
        p.add_run("Template Type: ").bold = True
        p.add_run(template.template_type)

    if template.subject_template:
        doc.add_heading("Subject", level=2)
        doc.add_paragraph(template.subject_template)

    doc.add_heading("Template Body", level=2)

    for line in (template.body_template or "").splitlines():
        clean_line = line.strip()

        if not clean_line:
            doc.add_paragraph("")
            continue

        if clean_line.startswith("•"):
            doc.add_paragraph(clean_line.replace("•", "", 1).strip(), style="List Bullet")
            continue

        if re.match(r"^\d+[.)]\s+", clean_line):
            item = re.sub(r"^\d+[.)]\s+", "", clean_line).strip()
            doc.add_paragraph(item, style="List Number")
            continue

        if re.match(r"^[-*]\s+", clean_line):
            item = re.sub(r"^[-*]\s+", "", clean_line).strip()
            doc.add_paragraph(item, style="List Bullet")
            continue

        doc.add_paragraph(clean_line)

    doc.save(file_path)

    return file_path