import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from bs4 import BeautifulSoup
from docx import Document


APP_DIR = Path(__file__).resolve().parents[1]
MOM_TEMPLATE_DIR = APP_DIR / "template_files" / "mom"
OUTPUT_DIR = APP_DIR.parent / "outputs" / "moms"


SECTION_HEADINGS = [
    "Technical Background / Context",
    "Topics Discussed",
    "Trade-offs / Alternatives Considered",
    "Technical Decisions & Rationale",
    "Action Items (Engineering)",
    "Risks / Technical Debt",
    "Project Overview (Scope & Objectives)",
    "Roles & Responsibilities",
    "Key Milestones & Timeline",
    "Communication Plan",
    "Decisions Taken",
    "Action Items",
    "Risks / Dependencies",
    "Next Steps",
    "Closing Summary",
]


def ensure_dirs():
    MOM_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def safe_filename(name: str) -> str:
    name = str(name or "MOM").strip()
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name[:120] or "MOM"


def validate_mom_template_file(filename: str) -> Path:
    ensure_dirs()

    if not filename:
        raise FileNotFoundError("MOM template file name is required")

    base = MOM_TEMPLATE_DIR.resolve()
    file_path = (MOM_TEMPLATE_DIR / filename).resolve()

    if file_path.parent != base:
        raise ValueError("Invalid MOM template file path")

    if file_path.suffix.lower() != ".docx":
        raise ValueError("Only .docx MOM templates are allowed")

    if not file_path.exists() or not file_path.is_file():
        raise FileNotFoundError(f"MOM template file not found: {filename}")

    return file_path


def clean_text(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def is_template_instruction(text: str) -> bool:
    t = clean_text(text).lower()

    bad_patterns = [
        "replace bracketed placeholders",
        "how to use this template",
        "tip:",
        "copy a row",
        "add more topics",
        "briefly describe",
        "describe a technical topic",
        "describe an alternative approach",
        "state the technical decision",
        "describe a technical risk",
        "write a short, polished closing",
        "enter meeting title",
        "enter date",
        "enter room",
        "enter attendee",
        "enter your name",
    ]

    if any(x in t for x in bad_patterns):
        return True

    if re.fullmatch(r"\d+\.?", t):
        return True

    if re.fullmatch(r"\[[^\]]+\]", t):
        return True

    return False


def split_lines(value: Any) -> List[str]:
    if value is None:
        return []

    if isinstance(value, list):
        return [clean_text(x).strip(" -•\t") for x in value if clean_text(x)]

    text = str(value or "").strip()
    if not text:
        return []

    text = BeautifulSoup(text, "html.parser").get_text("\n")
    parts = re.split(r"\n|;|\|", text)

    lines = []
    for part in parts:
        item = clean_text(part).strip(" -•\t")
        if item and not is_template_instruction(item):
            lines.append(item)

    return lines


def normalize_heading(text: str) -> str:
    text = clean_text(text)
    text = text.replace("■", "").replace("●", "").replace("•", "")
    return text.strip().lower()


def is_heading(text: str) -> bool:
    normalized = normalize_heading(text)
    return any(normalized == h.lower() for h in SECTION_HEADINGS)


def delete_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def delete_table(table):
    element = table._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def replace_paragraph_text(paragraph, replacements: Dict[str, str]):
    full_text = "".join(run.text for run in paragraph.runs)

    if not full_text:
        return

    new_text = full_text

    for old, new in replacements.items():
        new_text = new_text.replace(old, str(new or ""))

    if new_text == full_text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def replace_everywhere(doc: Document, replacements: Dict[str, str]):
    for paragraph in doc.paragraphs:
        replace_paragraph_text(paragraph, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph_text(paragraph, replacements)


def remove_bad_paragraphs(doc: Document):
    for paragraph in list(doc.paragraphs):
        text = clean_text(paragraph.text)
        if text and is_template_instruction(text):
            delete_paragraph(paragraph)


def clean_bracket_placeholders(doc: Document):
    pattern = re.compile(r"\[[^\]]+\]")

    for paragraph in doc.paragraphs:
        if pattern.search(paragraph.text):
            for run in paragraph.runs:
                run.text = pattern.sub("", run.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if pattern.search(paragraph.text):
                        for run in paragraph.runs:
                            run.text = pattern.sub("", run.text)


def remove_empty_numbered_paragraphs(doc: Document):
    for paragraph in list(doc.paragraphs):
        text = clean_text(paragraph.text)
        if not text or re.fullmatch(r"\d+\.?", text):
            delete_paragraph(paragraph)


def table_headers(table) -> List[str]:
    if not table.rows:
        return []
    return [clean_text(cell.text).lower() for cell in table.rows[0].cells]


def set_cell_text(cell, value: str):
    cell.text = clean_text(value)


def remove_extra_rows(table, keep_rows: int):
    while len(table.rows) > keep_rows:
        row = table.rows[-1]
        row._element.getparent().remove(row._element)


def extract_attendee_roles(attendees: str) -> List[Dict[str, str]]:
    people = []

    for item in re.split(r",|\n|;", attendees or ""):
        item = clean_text(item)
        if not item:
            continue

        role = ""
        match = re.search(r"\((.*?)\)", item)

        if match:
            role = clean_text(match.group(1)).lower()

        name = re.sub(r"\([^)]*\)", "", item).strip()

        people.append(
            {
                "name": name,
                "role": role,
                "full": item,
            }
        )

    return people


def get_person_by_role(people: List[Dict[str, str]], keywords: List[str]) -> Optional[str]:
    for person in people:
        role = person.get("role", "").lower()
        if any(keyword in role for keyword in keywords):
            return person.get("full")

    return None


def infer_owner(line: str, attendees: str) -> str:
    text = clean_text(line).lower()
    people = extract_attendee_roles(attendees)

    pm = get_person_by_role(
        people,
        [
            "project manager",
            "manager",
            "pm",
            "project coordinator",
            "coordinator",
            "scrum master",
        ],
    )

    tl = get_person_by_role(
        people,
        [
            "team lead",
            "technical lead",
            "tech lead",
            "tl",
            "lead",
            "technology owner",
            "technical owner",
            "architect",
        ],
    )

    dev = get_person_by_role(
        people,
        [
            "developer",
            "engineer",
            "software engineer",
            "frontend developer",
            "backend developer",
            "full stack",
            "programmer",
        ],
    )

    qa = get_person_by_role(
        people,
        [
            "qa",
            "tester",
            "quality",
            "test engineer",
        ],
    )

    business = get_person_by_role(
        people,
        [
            "business analyst",
            "ba",
            "client",
            "stakeholder",
            "product owner",
        ],
    )

    # Technical requirements should be owned by TL first.
    if any(
        x in text
        for x in [
            "technical requirement",
            "technical requirements",
            "technical planning",
            "technical discussion",
            "requirement",
            "requirements",
            "scope",
            "project workflow",
            "implementation plan",
            "planning",
            "client discussion",
            "business discussion",
            "resource introduction",
            "timeline",
            "milestone",
            "deadline",
            "approval",
        ]
    ):
        return tl or pm or business or dev or "To be assigned"

    # Frontend/backend implementation should be owned by developer.
    if any(
        x in text
        for x in [
            "frontend",
            "front-end",
            "react",
            "reactjs",
            "vite",
            "ui",
            "ux",
            "css",
            "html",
            "javascript",
            "typescript",
            "component",
            "page",
            "screen",
        ]
    ):
        return dev or tl or pm or "To be assigned"

    if any(
        x in text
        for x in [
            "backend",
            "back-end",
            "fastapi",
            "api",
            "database",
            "sql",
            "postgres",
            "mysql",
            "server",
            "endpoint",
            "authentication",
            "jwt",
        ]
    ):
        return dev or tl or pm or "To be assigned"

    if any(
        x in text
        for x in [
            "architecture",
            "technical direction",
            "technical design",
            "review",
            "code review",
            "standard",
            "deployment",
            "devops",
            "hosting",
            "infrastructure",
            "security",
        ]
    ):
        return tl or dev or pm or "To be assigned"

    if any(
        x in text
        for x in [
            "testing",
            "test",
            "qa",
            "bug",
            "quality",
            "validation",
        ]
    ):
        return qa or dev or tl or "To be assigned"

    return pm or tl or dev or business or "To be assigned"


def infer_priority(line: str, index: int = 0) -> str:
    text = line.lower()

    if any(x in text for x in ["urgent", "critical", "blocked", "client", "deadline", "high"]):
        return "High"

    if any(x in text for x in ["low", "routine", "later"]):
        return "Low"

    return "High" if index == 0 else "Medium"


def summarize_task(line: str) -> str:
    text = clean_text(line)
    lower = text.lower()

    if any(x in lower for x in ["requirement", "requirements", "scope", "workflow", "project"]):
        return "Finalize and document technical requirements"

    if any(x in lower for x in ["frontend", "front-end", "react", "reactjs", "vite", "ui", "ux"]):
        return "Prepare frontend implementation approach"

    if any(x in lower for x in ["backend", "back-end", "fastapi", "api", "database"]):
        return "Prepare backend API implementation approach"

    if any(x in lower for x in ["deployment", "hosting", "server", "devops"]):
        return "Plan deployment and hosting setup"

    if any(x in lower for x in ["testing", "qa", "bug", "validation"]):
        return "Define testing and validation plan"

    if len(text) > 90:
        return text[:87].rstrip() + "..."

    return text


def build_action_items(lines: List[str], attendees: str) -> List[Dict[str, str]]:
    action_items = []
    seen = set()

    for index, line in enumerate(lines):
        task = summarize_task(line)
        key = task.lower()

        if key in seen:
            continue

        seen.add(key)

        action_items.append(
            {
                "task": task,
                "owner": infer_owner(task, attendees),
                "deadline": "To be confirmed",
                "priority": infer_priority(line, index),
            }
        )

        if len(action_items) >= 5:
            break

    return action_items


def build_sections(raw_notes: str, attendees: str) -> Dict[str, Any]:
    lines = split_lines(raw_notes)

    if not lines:
        lines = ["Meeting discussion was completed and key points were documented."]

    text = " ".join(lines).lower()

    technical_lines = [
        line for line in lines
        if any(
            x in line.lower()
            for x in [
                "frontend",
                "front-end",
                "backend",
                "back-end",
                "api",
                "database",
                "react",
                "vite",
                "fastapi",
                "technical",
                "requirement",
                "requirements",
                "architecture",
                "deployment",
            ]
        )
    ]

    risk_lines = [
        line for line in lines
        if any(
            x in line.lower()
            for x in ["risk", "issue", "blocker", "dependency", "challenge", "debt"]
        )
    ]

    decision_lines = [
        line for line in lines
        if any(
            x in line.lower()
            for x in ["decided", "decision", "finalized", "approved", "agreed", "confirmed"]
        )
    ]

    topic_lines = technical_lines or lines

    if not decision_lines and any(
        x in text
        for x in ["frontend", "backend", "fastapi", "react", "vite", "technical", "requirement"]
    ):
        decision_lines = [
            "The team agreed on the discussed technical direction and implementation approach."
        ]

    action_items = build_action_items(topic_lines, attendees)

    return {
        "overview": " ".join(lines),
        "technical_background": (
            "The meeting focused on the technical requirements, implementation approach, "
            "tools, responsibilities, and next steps for the project."
        ),
        "topics": topic_lines,
        "tradeoffs": [],
        "decisions": decision_lines,
        "actions": action_items,
        "risks": risk_lines,
        "next_steps": [item["task"] for item in action_items[:3]],
        "closing": (
            "The meeting concluded with alignment on the discussed requirements, "
            "responsibilities, and next implementation steps."
        ),
    }


def remove_unused_sections(doc: Document, sections: Dict[str, Any]):
    remove_if_empty = {
        "Trade-offs / Alternatives Considered": sections.get("tradeoffs"),
        "Technical Decisions & Rationale": sections.get("decisions"),
        "Risks / Technical Debt": sections.get("risks"),
        "Risks / Dependencies": sections.get("risks"),
    }

    paragraphs = list(doc.paragraphs)

    for heading, data in remove_if_empty.items():
        if data:
            continue

        removing = False

        for paragraph in paragraphs:
            text = clean_text(paragraph.text)

            if not removing and normalize_heading(text) == heading.lower():
                removing = True
                delete_paragraph(paragraph)
                continue

            if removing:
                if is_heading(text):
                    break
                delete_paragraph(paragraph)


def fill_tables(doc: Document, sections: Dict[str, Any], attendees: str):
    people = extract_attendee_roles(attendees)

    for table in doc.tables:
        headers = table_headers(table)

        if "name" in headers and "role" in headers and "responsibility" in headers:
            if not people:
                delete_table(table)
                continue

            remove_extra_rows(table, 1)

            for person in people:
                row = table.add_row().cells
                set_cell_text(row[0], person.get("name", ""))
                set_cell_text(row[1], person.get("role", "Participant").title() or "Participant")
                set_cell_text(row[2], "Contribute to discussion and complete assigned responsibilities.")

        elif "milestone" in headers and "target date" in headers and "owner" in headers:
            next_steps = sections.get("next_steps") or []

            if not next_steps:
                delete_table(table)
                continue

            remove_extra_rows(table, 1)

            for step in next_steps[:3]:
                row = table.add_row().cells
                set_cell_text(row[0], step)
                set_cell_text(row[1], "To be confirmed")
                set_cell_text(row[2], infer_owner(step, attendees))

        elif "#" in headers and "task" in headers and "owner" in headers and "deadline" in headers:
            actions = sections.get("actions") or []

            if not actions:
                delete_table(table)
                continue

            remove_extra_rows(table, 1)

            for i, item in enumerate(actions, start=1):
                row = table.add_row().cells
                values = [
                    str(i),
                    item.get("task", "To be confirmed"),
                    item.get("owner", "To be assigned"),
                    item.get("deadline", "To be confirmed"),
                    item.get("priority", "Medium"),
                ]

                for j, value in enumerate(values):
                    if j < len(row):
                        set_cell_text(row[j], value)


def build_replacements(
    meeting_title: str,
    attendees: str,
    prepared_by: str,
    meeting_date: str,
    sections: Dict[str, Any],
) -> Dict[str, str]:
    topics = sections.get("topics") or []
    decisions = sections.get("decisions") or []
    risks = sections.get("risks") or []
    next_steps = sections.get("next_steps") or []

    return {
        "{{MEETING_TITLE}}": meeting_title,
        "{{MEETING_DATE}}": meeting_date,
        "{{ATTENDEES}}": attendees,
        "{{PREPARED_BY}}": prepared_by,
        "{{PROJECT_OVERVIEW}}": sections.get("overview", ""),
        "{{COMMUNICATION_PLAN}}": "The team will coordinate through regular updates and agreed communication channels.",
        "{{DECISIONS_TAKEN}}": "\n".join(decisions),
        "{{RISKS_DEPENDENCIES}}": "\n".join(risks) if risks else "No major risks or dependencies were identified.",
        "{{NEXT_STEPS}}": "\n".join(next_steps),
        "{{CLOSING_SUMMARY}}": sections.get("closing", ""),

        "[Enter meeting title]": meeting_title,
        "[Enter date]": meeting_date,
        "[Enter start – end time]": "To be confirmed",
        "[Enter room or video call link]": "To be confirmed",
        "[Enter attendee names, comma-separated]": attendees,
        "[Enter your name]": prepared_by,

        "[Briefly describe the technical problem or context being discussed.]": sections.get("technical_background", ""),
        "[Describe a technical topic, design point, or architecture detail discussed.]": topics[0] if topics else "",
        "[Add more topics as separate numbered lines.]": "\n".join(topics[1:]) if len(topics) > 1 else "",
        "Decision: [State the technical decision and the reasoning behind it.]": (
            "Decision: " + decisions[0] if decisions else ""
        ),
        "[Describe an alternative approach considered and why it was or wasn't chosen.]": "",
        "[Describe a technical risk or debt item introduced or identified. If none, write: No major technical risks were identified.]": (
            risks[0] if risks else ""
        ),
        "[Write a short, polished closing paragraph on the technical direction agreed upon.]": sections.get("closing", ""),

        "[Summarize the project scope, goals, and the business outcome it's meant to achieve.]": sections.get("overview", ""),
        "[Describe how often the team will sync, which channels will be used, and who to escalate to.]": (
            "The team will coordinate through regular updates and agreed communication channels."
        ),
        "Decision: [State the decision clearly and its business impact.]": (
            "Decision: " + decisions[0] if decisions else ""
        ),
        "[List any risk or dependency. If none were identified, write: No major risks or dependencies were identified.]": (
            risks[0] if risks else ""
        ),
        "[Write the immediate next step and who owns moving it forward.]": next_steps[0] if next_steps else "",
        "[Write a short, polished closing paragraph confirming the team is aligned to start.]": sections.get("closing", ""),
    }


def final_cleanup(doc: Document):
    remove_bad_paragraphs(doc)
    clean_bracket_placeholders(doc)
    remove_empty_numbered_paragraphs(doc)

    for table in list(doc.tables):
        all_text = " ".join(
            clean_text(cell.text)
            for row in table.rows
            for cell in row.cells
        )

        if not clean_text(all_text):
            delete_table(table)


def build_basic_mom_docx_from_template(
    template_filename: str,
    meeting_title: str,
    attendees: str,
    raw_notes: str,
    prepared_by: str,
    meeting_date: str | None = None,
) -> str:
    ensure_dirs()

    template_path = validate_mom_template_file(template_filename)
    doc = Document(str(template_path))

    meeting_date = meeting_date or datetime.now().strftime("%d %b %Y")
    sections = build_sections(raw_notes, attendees)

    replacements = build_replacements(
        meeting_title=meeting_title,
        attendees=attendees,
        prepared_by=prepared_by,
        meeting_date=meeting_date,
        sections=sections,
    )

    replace_everywhere(doc, replacements)
    remove_unused_sections(doc, sections)
    fill_tables(doc, sections, attendees)
    final_cleanup(doc)

    output_name = f"MOM - {safe_filename(meeting_title)} - {uuid.uuid4().hex[:8]}.docx"
    output_path = OUTPUT_DIR / output_name

    doc.save(str(output_path))
    return str(output_path)


def build_mom_docx_from_template(
    template_filename: str,
    meeting_title: str,
    attendees: str,
    prepared_by: str,
    meeting_date: str | None = None,
    meeting_objective: str = "",
    meeting_overview: str = "",
    key_discussion_points: Any = None,
    decisions_taken: Any = None,
    action_items: Any = None,
    risks_dependencies: Any = None,
    next_steps: Any = None,
    closing_summary: str = "",
    generated_mom_text: str = "",
) -> str:
    raw_notes = (
        generated_mom_text
        or meeting_overview
        or meeting_objective
        or ""
    )

    return build_basic_mom_docx_from_template(
        template_filename=template_filename,
        meeting_title=meeting_title,
        attendees=attendees,
        raw_notes=raw_notes,
        prepared_by=prepared_by,
        meeting_date=meeting_date,
    )