from pathlib import Path
from html import escape

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import FileResponse
from docx import Document


router = APIRouter(
    prefix="/email-template-files",
    tags=["Email Template Files"],
)

APP_DIR = Path(__file__).resolve().parents[1]
BASE_DIR = APP_DIR / "template_files" / "email"


def validate_category(category: str) -> Path:
    if not category:
        raise HTTPException(status_code=400, detail="Category is required")

    base = BASE_DIR.resolve()
    folder = (BASE_DIR / category).resolve()

    if base != folder and base not in folder.parents:
        raise HTTPException(status_code=400, detail="Invalid category")

    if not folder.exists() or not folder.is_dir():
        raise HTTPException(status_code=404, detail=f"Category not found: {folder}")

    return folder


def validate_template_file(category: str, filename: str) -> Path:
    if not filename:
        raise HTTPException(status_code=400, detail="Template file is required")

    folder = validate_category(category)
    file_path = (folder / filename).resolve()

    if folder.resolve() not in file_path.parents:
        raise HTTPException(status_code=400, detail="Invalid template file")

    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(status_code=404, detail=f"Template file not found: {file_path}")

    if file_path.suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Only DOCX templates are allowed")

    return file_path


def paragraph_has_numbering(paragraph) -> bool:
    try:
        p_pr = paragraph._p.pPr
        return p_pr is not None and p_pr.numPr is not None
    except Exception:
        return False


def get_paragraph_style_name(paragraph) -> str:
    try:
        return (paragraph.style.name or "").lower()
    except Exception:
        return ""


def read_runs_as_html(paragraph) -> str:
    parts = []

    for run in paragraph.runs:
        text = run.text

        if not text:
            continue

        safe_text = escape(text)

        if run.bold:
            parts.append(f"<strong>{safe_text}</strong>")
        else:
            parts.append(safe_text)

    if parts:
        return "".join(parts).strip()

    return escape(paragraph.text.strip())


def get_docx_paragraph_html(paragraph) -> dict | None:
    text = paragraph.text.strip()

    if not text:
        return None

    style_name = get_paragraph_style_name(paragraph)
    has_numbering = paragraph_has_numbering(paragraph)
    html = read_runs_as_html(paragraph)

    if "bullet" in style_name or has_numbering:
        return {
            "type": "bullet",
            "html": html,
            "text": f"• {text}",
        }

    if "number" in style_name:
        return {
            "type": "number",
            "html": html,
            "text": f"1. {text}",
        }

    return {
        "type": "paragraph",
        "html": html,
        "text": text,
    }


def read_docx_text(file_path: Path) -> str:
    document = Document(str(file_path))
    parts = []

    for paragraph in document.paragraphs:
        item = get_docx_paragraph_html(paragraph)

        if not item:
            continue

        if item["type"] == "bullet":
            parts.append(f"• {item['html']}")
        elif item["type"] == "number":
            parts.append(f"1. {item['html']}")
        else:
            parts.append(item["html"])

    for table in document.tables:
        for row in table.rows:
            row_items = []

            for cell in row.cells:
                cell_lines = []

                for paragraph in cell.paragraphs:
                    item = get_docx_paragraph_html(paragraph)

                    if not item:
                        continue

                    if item["type"] == "bullet":
                        cell_lines.append(f"• {item['html']}")
                    elif item["type"] == "number":
                        cell_lines.append(f"1. {item['html']}")
                    else:
                        cell_lines.append(item["html"])

                cell_text = " ".join(cell_lines).strip()

                if cell_text:
                    row_items.append(cell_text)

            if row_items:
                parts.append(" | ".join(row_items))

    return "\n".join(parts).strip()


@router.get("/debug")
def debug_email_template_path():
    return {
        "app_dir": str(APP_DIR),
        "base_dir": str(BASE_DIR),
        "base_dir_exists": BASE_DIR.exists(),
        "items": [item.name for item in BASE_DIR.iterdir()] if BASE_DIR.exists() else [],
    }


@router.get("/categories")
def get_email_template_categories():
    if not BASE_DIR.exists():
        return {
            "categories": [],
            "error": f"Folder not found: {BASE_DIR}",
        }

    categories = []

    for folder in sorted(BASE_DIR.iterdir(), key=lambda x: x.name.lower()):
        if folder.is_dir():
            docx_files = list(folder.rglob("*.docx"))
            categories.append(
                {
                    "name": folder.name,
                    "count": len(docx_files),
                }
            )

    return {"categories": categories}


@router.get("/subcategories")
def get_email_template_subcategories(category: str = Query(...)):
    folder = validate_category(category)

    templates = []

    for file in sorted(folder.rglob("*.docx"), key=lambda x: x.name.lower()):
        relative_file = file.relative_to(folder)
        templates.append(
            {
                "name": file.stem,
                "filename": str(relative_file).replace("\\", "/"),
            }
        )

    return {"templates": templates}


@router.get("/content")
def get_email_template_content(
    category: str = Query(...),
    filename: str = Query(...),
):
    file_path = validate_template_file(category, filename)

    return {
        "category": category,
        "template_name": file_path.stem,
        "filename": filename,
        "content": read_docx_text(file_path),
    }


@router.get("/download")
def download_email_template_file(
    category: str = Query(...),
    filename: str = Query(...),
):
    file_path = validate_template_file(category, filename)

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )