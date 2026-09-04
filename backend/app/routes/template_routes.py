import os
import re
import tempfile
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from docx import Document

from app import models, schemas
from app.auth import get_current_user
from app.config import settings
from app.database import get_db

router = APIRouter(prefix="/templates", tags=["Templates"])


def safe_filename(name: str) -> str:
    name = name or "template"
    name = re.sub(r"[^\w\s-]", "", name).strip()
    name = re.sub(r"\s+", "_", name)
    return name[:80] or "template"


def resolve_template_file_path(source_file: str | None) -> Path | None:
    if not source_file:
        return None

    raw_path = Path(source_file)

    possible_paths = []

    if raw_path.is_absolute():
        possible_paths.append(raw_path)
    else:
        possible_paths.extend(
            [
                raw_path,
                Path(settings.OUTPUT_DIR) / "templates" / source_file,
                Path(settings.OUTPUT_DIR) / "templates" / "email" / source_file,
                Path(settings.OUTPUT_DIR) / source_file,
                Path("outputs") / "templates" / source_file,
                Path("outputs") / "templates" / "email" / source_file,
                Path("templates") / source_file,
                Path("templates") / "email" / source_file,
                Path("app") / "templates" / source_file,
                Path("app") / "templates" / "email" / source_file,
                Path("app") / "data" / source_file,
            ]
        )

    for path in possible_paths:
        if path.exists() and path.is_file():
            return path

    return None


def create_docx_from_template_library(template: models.TemplateLibrary) -> str:
    document = Document()

    document.add_heading(template.template_name or "Email Template", level=1)

    if template.category:
        p = document.add_paragraph()
        p.add_run("Category: ").bold = True
        p.add_run(template.category)

    if template.subject_template:
        p = document.add_paragraph()
        p.add_run("Subject: ").bold = True
        p.add_run(template.subject_template)

    document.add_paragraph("")

    for line in (template.body_template or "").splitlines():
        document.add_paragraph(line.strip() if line.strip() else "")

    output_path = os.path.join(
        tempfile.gettempdir(),
        f"{safe_filename(template.template_name)}.docx",
    )

    document.save(output_path)
    return output_path


@router.get("")
def list_templates(
    type: str | None = Query(None),
    template_type: str = Query("meeting_mom"),
    category: str | None = Query(None),
    search: str | None = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if type == "email":
        query = db.query(models.TemplateLibrary).filter(
            models.TemplateLibrary.template_type == "email",
            models.TemplateLibrary.is_active == True,
        )

        if category:
            query = query.filter(models.TemplateLibrary.category == category)

        if search:
            search_term = f"%{search}%"
            query = query.filter(models.TemplateLibrary.template_name.ilike(search_term))

        return (
            query.order_by(
                models.TemplateLibrary.category.asc(),
                models.TemplateLibrary.template_name.asc(),
            )
            .all()
        )

    return (
        db.query(models.CommunicationTemplate)
        .filter(
            models.CommunicationTemplate.user_id == current_user.id,
            models.CommunicationTemplate.template_type == template_type,
        )
        .order_by(models.CommunicationTemplate.created_at.desc())
        .all()
    )


@router.get("/{template_id}/download")
def download_template_docx(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    template = (
        db.query(models.TemplateLibrary)
        .filter(
            models.TemplateLibrary.id == template_id,
            models.TemplateLibrary.template_type == "email",
            models.TemplateLibrary.is_active == True,
        )
        .first()
    )

    if not template:
        raise HTTPException(status_code=404, detail="Email template not found")

    source_path = resolve_template_file_path(template.source_file)

    if source_path:
        return FileResponse(
            path=str(source_path),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=f"{safe_filename(template.template_name)}.docx",
        )

    generated_docx_path = create_docx_from_template_library(template)

    return FileResponse(
        path=generated_docx_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{safe_filename(template.template_name)}.docx",
    )


@router.get("/{template_id}")
def get_template(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    library_template = (
        db.query(models.TemplateLibrary)
        .filter(
            models.TemplateLibrary.id == template_id,
            models.TemplateLibrary.template_type == "email",
            models.TemplateLibrary.is_active == True,
        )
        .first()
    )

    if library_template:
        return library_template

    template = (
        db.query(models.CommunicationTemplate)
        .filter(
            models.CommunicationTemplate.id == template_id,
            models.CommunicationTemplate.user_id == current_user.id,
        )
        .first()
    )

    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    return template


@router.post("", response_model=schemas.TemplateOut, status_code=201)
def create_template(
    req: schemas.TemplateCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    template = models.CommunicationTemplate(
        user_id=current_user.id,
        template_type=req.template_type,
        name=req.name.strip(),
        description=req.description,
        meeting_title=req.meeting_title,
        attendees=req.attendees,
        raw_notes=req.raw_notes,
        template_content=req.template_content,
    )

    db.add(template)
    db.commit()
    db.refresh(template)

    return template


@router.put("/{template_id}", response_model=schemas.TemplateOut)
def update_template(
    template_id: int,
    req: schemas.TemplateUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    template = (
        db.query(models.CommunicationTemplate)
        .filter(
            models.CommunicationTemplate.id == template_id,
            models.CommunicationTemplate.user_id == current_user.id,
        )
        .first()
    )

    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    data = req.model_dump(exclude_unset=True)

    for key, value in data.items():
        setattr(template, key, value)

    db.commit()
    db.refresh(template)

    return template


@router.delete("/{template_id}")
def delete_template(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    template = (
        db.query(models.CommunicationTemplate)
        .filter(
            models.CommunicationTemplate.id == template_id,
            models.CommunicationTemplate.user_id == current_user.id,
        )
        .first()
    )

    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    db.delete(template)
    db.commit()

    return {"message": "Template deleted successfully"}