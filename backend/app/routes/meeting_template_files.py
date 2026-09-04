from pathlib import Path

from docx import Document
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import FileResponse

router = APIRouter(
    prefix="/meeting-template-files",
    tags=["Meeting Template Files"],
)

APP_DIR = Path(__file__).resolve().parents[1]
BASE_DIR = APP_DIR / "template_files" / "mom"


def validate_template_file(filename: str) -> Path:
    if not filename:
        raise HTTPException(status_code=400, detail="Template file is required")

    BASE_DIR.mkdir(parents=True, exist_ok=True)

    base = BASE_DIR.resolve()
    file_path = (BASE_DIR / filename).resolve()

    if file_path.parent != base:
        raise HTTPException(status_code=400, detail="Invalid template file path")

    if file_path.suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="Only .docx templates are allowed")

    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(status_code=404, detail=f"Template file not found: {filename}")

    return file_path


@router.get("/debug")
def debug_mom_templates():
    BASE_DIR.mkdir(parents=True, exist_ok=True)

    return {
        "base_dir": str(BASE_DIR),
        "exists": BASE_DIR.exists(),
        "files": [item.name for item in BASE_DIR.iterdir()] if BASE_DIR.exists() else [],
    }


@router.get("/categories")
def list_mom_categories():
    BASE_DIR.mkdir(parents=True, exist_ok=True)

    files = [
        item.name
        for item in BASE_DIR.iterdir()
        if item.is_file()
        and item.suffix.lower() == ".docx"
        and not item.name.startswith("~$")
        and not item.name.startswith(".")
    ]

    return {"categories": sorted(files, key=str.lower)}


@router.get("/files")
def list_mom_template_files():
    data = list_mom_categories()
    return {"files": data["categories"]}


@router.get("/download")
def download_mom_template_file(filename: str = Query(...)):
    file_path = validate_template_file(filename)

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@router.get("/content")
def read_mom_template_content(filename: str = Query(...)):
    file_path = validate_template_file(filename)

    try:
        doc = Document(str(file_path))
        lines = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_items = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        row_items.append(text)
                if row_items:
                    lines.append(" | ".join(row_items))

        return {
            "filename": filename,
            "content": "\n".join(lines).strip(),
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Unable to read DOCX template: {str(e)}",
        )